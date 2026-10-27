"""เซิร์ฟเวอร์บอร์ดงาน: ใช้ taskjob.xlsx เป็นฐานข้อมูล

รัน:  .venv/bin/python server.py   แล้วเปิด http://localhost:8765
"""
import json
import os
import re
import tempfile
import threading
import time
import uuid
from datetime import date, datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(ROOT, "taskjob.xlsx")
SHEET = "Tasks"
PORT = int(os.environ.get("PORT", "8765"))

# (field, header in Excel, column width)
COLUMNS = [
    ("id", "ID", 18),
    ("key", "รหัส", 10),
    ("title", "ชื่องาน", 40),
    ("status", "สถานะ", 12),
    ("type", "ประเภท", 10),
    ("priority", "ความสำคัญ", 12),
    ("assignee", "ผู้รับผิดชอบ", 18),
    ("points", "Story points", 12),
    ("due", "กำหนดส่ง", 13),
    ("desc", "รายละเอียด", 50),
    ("order", "ลำดับ", 8),
    ("created", "สร้างเมื่อ", 18),
    ("updated", "อัปเดตเมื่อ", 18),
]

# Excel shows Thai labels; the web page uses the codes.
STATUS = {"todo": "ต้องทำ", "doing": "กำลังทำ", "review": "รอตรวจ", "done": "เสร็จแล้ว"}
TYPE = {"task": "Task", "bug": "Bug", "story": "Story"}
PRIORITY = {"highest": "สูงสุด", "high": "สูง", "medium": "ปานกลาง", "low": "ต่ำ"}

lock = threading.Lock()


def to_code(value, table, default):
    """Accept either a code ("todo") or its Excel label ("ต้องทำ")."""
    v = str(value or "").strip()
    if v.lower() in table:
        return v.lower()
    for code, label in table.items():
        if v == label or v.lower() == label.lower():
            return code
    return default


def to_ms(value):
    if isinstance(value, datetime):
        return int(value.timestamp() * 1000)
    if isinstance(value, (int, float)):
        return int(value)
    return None


def to_date_str(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    v = str(value or "").strip()
    return v if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) else ""


def to_int_key(value):
    m = re.search(r"\d+", str(value or ""))
    return int(m.group()) if m else None


def new_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    return wb


def prepare_sheet(wb):
    """Make sure the Tasks sheet exists with a styled header row and dropdowns."""
    if SHEET in wb.sheetnames:
        ws = wb[SHEET]
    elif len(wb.sheetnames) == 1 and wb.active.max_row <= 1 and wb.active["A1"].value is None:
        ws = wb.active  # the empty Sheet1 the file started with
        ws.title = SHEET
    else:
        ws = wb.create_sheet(SHEET, 0)

    if ws["A1"].value != COLUMNS[0][1]:
        head_fill = PatternFill("solid", fgColor="1F5FBF")
        for i, (_, header, width) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=1, column=i, value=header)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
            ws.column_dimensions[c.column_letter].width = width
        ws.freeze_panes = "A2"
        for field, table in (("status", STATUS), ("type", TYPE), ("priority", PRIORITY)):
            col = [f for f, _, _ in COLUMNS].index(field) + 1
            letter = ws.cell(row=1, column=col).column_letter
            dv = DataValidation(type="list", formula1='"' + ",".join(table.values()) + '"', allow_blank=True)
            dv.add(f"{letter}2:{letter}2000")
            ws.add_data_validation(dv)
    return ws


def header_map(ws):
    """field -> column index, matched by header text so columns can be reordered."""
    by_header = {h: f for f, h, _ in COLUMNS}
    found = {}
    for cell in ws[1]:
        f = by_header.get(str(cell.value or "").strip())
        if f:
            found[f] = cell.column
    return found


def open_book():
    if os.path.exists(XLSX) and os.path.getsize(XLSX) > 0:
        wb = load_workbook(XLSX)
    else:
        wb = new_workbook()
    return wb, prepare_sheet(wb)


def read_tasks(ws):
    cols = header_map(ws)
    tasks, needs_fix = [], False
    for row in ws.iter_rows(min_row=2):
        get = lambda f: row[cols[f] - 1].value if f in cols and cols[f] - 1 < len(row) else None
        title = str(get("title") or "").strip()
        if not title:
            continue
        t = {
            "id": str(get("id") or "").strip(),
            "key": to_int_key(get("key")),
            "title": title,
            "status": to_code(get("status"), STATUS, "todo"),
            "type": to_code(get("type"), TYPE, "task"),
            "priority": to_code(get("priority"), PRIORITY, "medium"),
            "assignee": str(get("assignee") or "").strip(),
            "points": int(get("points")) if isinstance(get("points"), (int, float)) else 0,
            "due": to_date_str(get("due")),
            "desc": str(get("desc") or ""),
            "order": get("order") if isinstance(get("order"), (int, float)) else None,
            "created": to_ms(get("created")),
            "updated": to_ms(get("updated")),
        }
        tasks.append(t)

    # Rows typed straight into Excel may lack an ID, key or order: fill them in.
    now = int(time.time() * 1000)
    max_key = max([t["key"] or 0 for t in tasks] + [0])
    seen = set()
    for t in tasks:
        if not t["id"] or t["id"] in seen:
            t["id"] = "t" + uuid.uuid4().hex[:10]
            needs_fix = True
        seen.add(t["id"])
        if not t["key"]:
            max_key += 1
            t["key"] = max_key
            needs_fix = True
        if t["order"] is None:
            same = [x["order"] for x in tasks if x["status"] == t["status"] and x["order"] is not None]
            t["order"] = (max(same) if same else 0) + 1
            needs_fix = True
        if t["created"] is None:
            t["created"] = now
            needs_fix = True
        if t["updated"] is None:
            t["updated"] = t["created"]
            needs_fix = True
    return tasks, needs_fix


def write_tasks(wb, ws, tasks):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    cols = header_map(ws)
    for r, t in enumerate(sorted(tasks, key=lambda t: t["key"]), start=2):
        values = {
            "id": t["id"],
            "key": f"TSK-{t['key']}",
            "title": t["title"],
            "status": STATUS[t["status"]],
            "type": TYPE[t["type"]],
            "priority": PRIORITY[t["priority"]],
            "assignee": t["assignee"] or None,
            "points": t["points"] or None,
            "due": date.fromisoformat(t["due"]) if t["due"] else None,
            "desc": t["desc"] or None,
            "order": t["order"],
            "created": datetime.fromtimestamp(t["created"] / 1000),
            "updated": datetime.fromtimestamp(t["updated"] / 1000),
        }
        for f, col in cols.items():
            c = ws.cell(row=r, column=col, value=values[f])
            if f == "due" and values[f]:
                c.number_format = "yyyy-mm-dd"
            elif f in ("created", "updated"):
                c.number_format = "yyyy-mm-dd hh:mm"
            elif f == "desc":
                c.alignment = Alignment(wrap_text=True, vertical="top")
    # Write to a temp file first so a crash never leaves a half-written workbook.
    fd, tmp = tempfile.mkstemp(dir=ROOT, suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(tmp)
        os.chmod(tmp, 0o644)  # mkstemp creates 0600; keep the workbook readable like a normal file
        os.replace(tmp, XLSX)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def clean(data, existing, tasks):
    """Validate a task sent from the page."""
    title = str(data.get("title") or "").strip()
    if not title:
        raise ValueError("ต้องมีชื่องาน")
    now = int(time.time() * 1000)
    key = to_int_key(data.get("key")) or (existing or {}).get("key")
    taken = {t["key"] for t in tasks if not existing or t["id"] != existing["id"]}
    if not key or key in taken:
        key = max(taken | {0}) + 1
    try:
        points = max(0, int(data.get("points") or 0))
    except (TypeError, ValueError):
        points = 0
    order = data.get("order")
    return {
        "key": key,
        "title": title[:500],
        "status": to_code(data.get("status"), STATUS, "todo"),
        "type": to_code(data.get("type"), TYPE, "task"),
        "priority": to_code(data.get("priority"), PRIORITY, "medium"),
        "assignee": str(data.get("assignee") or "").strip()[:100],
        "points": points,
        "due": to_date_str(data.get("due")),
        "desc": str(data.get("desc") or "")[:10000],
        "order": order if isinstance(order, (int, float)) else 1,
        "created": to_ms(data.get("created")) or (existing or {}).get("created") or now,
        "updated": now,
    }


# ---------- notes: one Markdown file per note in ./document ----------
NOTES_DIR = os.path.join(ROOT, "document")
TRASH_DIR = os.path.join(NOTES_DIR, ".trash")
NOTE_EXT = (".md", ".txt")
notes_lock = threading.Lock()


def safe_filename(title):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "-", title).strip().lstrip(".")
    return re.sub(r"\s+", " ", name)[:80].strip() or "ไม่มีชื่อ"


def note_path(note_id):
    """Resolve an id (the file name) to a path inside NOTES_DIR, or None if it isn't one."""
    if not note_id or os.path.basename(note_id) != note_id or note_id.startswith(".") or not note_id.endswith(NOTE_EXT):
        return None
    path = os.path.join(NOTES_DIR, note_id)
    return path if os.path.isfile(path) else None


def read_note(note_id, with_body=True):
    path = os.path.join(NOTES_DIR, note_id)
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    title, body = os.path.splitext(note_id)[0], text
    first, _, rest = text.partition("\n")
    if first.startswith("# "):
        title, body = first[2:].strip() or title, rest.lstrip("\n")
    note = {"id": note_id, "title": title, "updated": int(os.path.getmtime(path) * 1000)}
    if with_body:
        note["body"] = body
    else:
        note["preview"] = re.sub(r"\s+", " ", body).strip()[:120]
    return note


def list_notes():
    os.makedirs(NOTES_DIR, exist_ok=True)
    notes = [read_note(n, with_body=False) for n in os.listdir(NOTES_DIR)
             if n.endswith(NOTE_EXT) and not n.startswith(".") and os.path.isfile(os.path.join(NOTES_DIR, n))]
    return sorted(notes, key=lambda n: n["updated"], reverse=True)


def write_note(old_id, title, body):
    """Save a note; the file name follows the title, so a new title renames the file."""
    os.makedirs(NOTES_DIR, exist_ok=True)
    title = re.sub(r"\s+", " ", str(title or "")).strip()[:200] or "ไม่มีชื่อ"
    body = str(body or "")[:500000]
    base = safe_filename(title)
    new_id, n = base + ".md", 2
    while new_id != old_id and os.path.exists(os.path.join(NOTES_DIR, new_id)):
        new_id, n = f"{base} ({n}).md", n + 1
    fd, tmp = tempfile.mkstemp(dir=NOTES_DIR, prefix=".tmp-", suffix=".md")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(f"# {title}\n\n{body}")
    os.chmod(tmp, 0o644)
    os.replace(tmp, os.path.join(NOTES_DIR, new_id))
    if old_id and old_id != new_id:
        os.remove(os.path.join(NOTES_DIR, old_id))
    return read_note(new_id)


def trash_note(note_id):
    """Deleted notes go to document/.trash so they can be recovered by hand."""
    os.makedirs(TRASH_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    os.replace(os.path.join(NOTES_DIR, note_id), os.path.join(TRASH_DIR, f"{stamp} {note_id}"))


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=ROOT, **kwargs)

    def log_message(self, fmt, *args):
        if "/api/" in self.path:
            super().log_message(fmt, *args)

    def send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def task_id(self):
        m = re.fullmatch(r"/api/tasks/([A-Za-z0-9_-]{1,64})", self.path)
        return m.group(1) if m else None

    def handle_api(self, change=None):
        try:
            with lock:
                wb, ws = open_book()
                tasks, needs_fix = read_tasks(ws)
                if change:
                    tasks = change(tasks)
                if change or needs_fix or not os.path.exists(XLSX) or os.path.getsize(XLSX) == 0:
                    write_tasks(wb, ws, tasks)
            self.send_json(200, {"tasks": tasks})
        except ValueError as e:
            self.send_json(400, {"error": str(e)})
        except PermissionError:
            self.send_json(423, {"error": "เขียนไฟล์ taskjob.xlsx ไม่ได้ ปิดไฟล์ใน Excel แล้วลองอีกครั้ง"})
        except Exception as e:  # noqa: BLE001
            self.send_json(500, {"error": f"อ่านหรือเขียน Excel ไม่สำเร็จ: {e}"})

    def read_json(self):
        length = int(self.headers.get("Content-Length") or 0)
        data = json.loads(self.rfile.read(length) or b"{}")
        if not isinstance(data, dict):
            raise ValueError
        return data

    def notes_api(self, method):
        rest = unquote(self.path.split("?")[0][len("/api/notes"):]).lstrip("/")
        try:
            with notes_lock:
                if not rest:
                    if method == "GET":
                        return self.send_json(200, {"notes": list_notes()})
                    if method == "POST":
                        data = self.read_json()
                        return self.send_json(200, {"note": write_note(None, data.get("title"), data.get("body"))})
                    return self.send_json(405, {"error": "method not allowed"})
                if not note_path(rest):
                    return self.send_json(404, {"error": "ไม่พบโน้ตนี้ อาจถูกลบหรือเปลี่ยนชื่อไปแล้ว"})
                if method == "GET":
                    return self.send_json(200, {"note": read_note(rest)})
                if method == "PUT":
                    data = self.read_json()
                    return self.send_json(200, {"note": write_note(rest, data.get("title"), data.get("body"))})
                if method == "DELETE":
                    trash_note(rest)
                    return self.send_json(200, {"ok": True})
                return self.send_json(405, {"error": "method not allowed"})
        except (ValueError, json.JSONDecodeError):
            self.send_json(400, {"error": "ข้อมูลไม่ถูกต้อง"})
        except OSError as e:
            self.send_json(500, {"error": f"บันทึกไฟล์ในโฟลเดอร์ document ไม่สำเร็จ: {e.strerror or e}"})

    def is_notes(self):
        return self.path == "/api/notes" or self.path.startswith(("/api/notes/", "/api/notes?"))

    def do_POST(self):
        if self.is_notes():
            return self.notes_api("POST")
        self.send_json(404, {"error": "not found"})

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            self.path = "/taskboard.html"
        if self.is_notes():
            return self.notes_api("GET")
        if self.path == "/api/tasks":
            return self.handle_api()
        if self.path.startswith("/api/"):
            return self.send_json(404, {"error": "not found"})
        if not self.path.split("?")[0].endswith((".html", ".js")):
            return self.send_json(404, {"error": "not found"})  # serve only the pages and their scripts, never the workbook or server code
        return super().do_GET()

    def do_HEAD(self):
        self.send_json(404, {"error": "not found"})

    def end_headers(self):
        if not self.path.startswith("/api/"):
            self.send_header("Cache-Control", "no-store")  # always serve the latest page
        super().end_headers()

    def do_PUT(self):
        if self.is_notes():
            return self.notes_api("PUT")
        tid = self.task_id()
        if not tid:
            return self.send_json(404, {"error": "not found"})
        try:
            length = int(self.headers.get("Content-Length") or 0)
            data = json.loads(self.rfile.read(length) or b"{}")
        except (ValueError, json.JSONDecodeError):
            return self.send_json(400, {"error": "ข้อมูลไม่ถูกต้อง"})

        def change(tasks):
            existing = next((t for t in tasks if t["id"] == tid), None)
            t = {"id": tid, **clean(data, existing, tasks)}
            return [t if x["id"] == tid else x for x in tasks] if existing else tasks + [t]

        self.handle_api(change)

    def do_DELETE(self):
        if self.is_notes():
            return self.notes_api("DELETE")
        tid = self.task_id()
        if not tid:
            return self.send_json(404, {"error": "not found"})
        self.handle_api(lambda tasks: [t for t in tasks if t["id"] != tid])


if __name__ == "__main__":
    with lock:
        wb, ws = open_book()
        tasks, _ = read_tasks(ws)
        write_tasks(wb, ws, tasks)  # creates the header row on first run
    os.makedirs(NOTES_DIR, exist_ok=True)
    print(f"บอร์ดงานพร้อมแล้ว: http://localhost:{PORT}")
    print(f"ข้อมูลเก็บที่: {XLSX}")
    print(f"โน้ตเก็บที่: {NOTES_DIR}")
    print("กด Ctrl+C เพื่อหยุด")
    ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()
